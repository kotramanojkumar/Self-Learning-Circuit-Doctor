"""
generate_presentation.py - Builds the 16-Slide Professional PPTX Presentation
Part of 'Self-Learning Circuit Doctor: Voltage-Signature-Based Fault Detection in DC Circuits'
BEEE Project for B.Tech
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(output_path: str):
    prs = Presentation()
    # Set slide dimensions to widescreen 16:9 (13.33 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6] # Blank layout

    # Professional Color Palette
    PRIMARY_BLUE = RGBColor(26, 54, 93)      # #1A365D Deep Navy
    SECONDARY_BLUE = RGBColor(49, 130, 206)  # #3182CE Accent Blue
    TEXT_DARK = RGBColor(45, 55, 72)         # #2D3748 Charcoal Body Text
    TEXT_MUTED = RGBColor(113, 128, 150)     # #718096 Slate Gray
    BG_LIGHT = RGBColor(247, 250, 252)       # #F7FAFC Soft background
    WHITE = RGBColor(255, 255, 255)
    ACCENT_GREEN = RGBColor(56, 161, 105)    # #38A169 Emerald
    CARD_BG = RGBColor(237, 242, 247)        # #EDF2F7 Card Gray

    base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
    fig_dir = os.path.join(base_dir, 'figures')

    def add_header(slide, title_text, category="B.TECH BEEE PROJECT"):
        # Header banner
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_cat = tf.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = SECONDARY_BLUE

        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = PRIMARY_BLUE

    def add_speaker_note(slide, note_text):
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = note_text

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    # Background card
    bg = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = PRIMARY_BLUE
    bg.line.fill.background()

    tb = slide1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.333), Inches(3.2))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "B.TECH PROJECT (BEEE)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = SECONDARY_BLUE

    p2 = tf.add_paragraph()
    p2.text = "Self-Learning Circuit Doctor:\nVoltage-Signature-Based Fault Detection in DC Circuits"
    p2.font.size = Pt(32)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_before = Pt(12)

    p3 = tf.add_paragraph()
    p3.text = "Online Circuit Simulation (Falstad) & Machine-Learning-Driven Diagnostic Framework"
    p3.font.size = Pt(16)
    p3.font.color.rgb = RGBColor(226, 232, 240)
    p3.space_before = Pt(8)

    # Team Box
    team_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.8), Inches(11.333), Inches(1.8))
    team_box.fill.solid()
    team_box.fill.fore_color.rgb = RGBColor(44, 82, 130)
    team_box.line.fill.background()

    ttb = slide1.shapes.add_textbox(Inches(1.2), Inches(4.9), Inches(10.9), Inches(1.6))
    ttf = ttb.text_frame
    ttf.word_wrap = True

    tp1 = ttf.paragraphs[0]
    tp1.text = "PROJECT TEAM (TWO-MEMBER COLLABORATION)"
    tp1.font.size = Pt(11)
    tp1.font.bold = True
    tp1.font.color.rgb = SECONDARY_BLUE

    tp2 = ttf.add_paragraph()
    tp2.text = "• Student Member 1: KOTRA MANOJ KUMAR (Reg. No: 23BCE2113) — CSE (SCOPE)\n" \
               "• Student Member 2: Mayank Yadav (Reg. No: 23BCE0119) — CSE (SCOPE)\n" \
               "Faculty Supervisor & Guide: Prof. POOMPAVAI T | School of Electrical Engineering (SELECT) | VIT"
    tp2.font.size = Pt(12)
    tp2.font.color.rgb = WHITE
    tp2.space_before = Pt(4)

    add_speaker_note(slide1,
        "Good morning respected evaluators and faculty members. Welcome to our BEEE project presentation. "
        "Our project is titled 'Self-Learning Circuit Doctor: Voltage-Signature-Based Fault Detection in DC Circuits'. "
        "I am Member 1, focusing on electrical modeling and Falstad circuit simulation, and my co-presenter is Member 2, "
        "who spearheaded the Python diagnostic algorithms and graphical user interface. "
        "In this presentation, we will demonstrate how fundamental circuit laws like KCL, KVL, and Ohm's law can be combined "
        "with multi-signal signature pattern recognition to diagnose and localize 13 distinct operating conditions in DC networks with 100% precision."
    )

    # =========================================================================
    # SLIDE 2: Problem Statement & Motivation
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Problem Statement & Engineering Motivation")

    # Card 1: The Challenge
    c1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = RGBColor(203, 213, 225)
    tb1 = slide2.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "The Industrial & Educational Challenge"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    bullets1 = [
        "Component Degradation: Resistors in industrial power supplies and automotive ECUs suffer from thermal cracking, vibrational open circuits, solder bridges, and value drift.",
        "Manual Troubleshooting Bottlenecks: Technicians spend hours desoldering parts or probing dozens of test points sequentially with multimeters.",
        "Educational Gap: Engineering students often study KCL and nodal analysis purely as abstract mathematical exercises without seeing how they automate fault diagnosis."
    ]
    for b in bullets1:
        p_b = tf1.add_paragraph()
        p_b.text = "• " + b
        p_b.font.size = Pt(12)
        p_b.font.color.rgb = TEXT_DARK
        p_b.space_before = Pt(8)

    # Card 2: The Proposed Solution
    c2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = RGBColor(235, 248, 255)
    c2.line.color.rgb = SECONDARY_BLUE
    tb2 = slide2.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "The 'Circuit Doctor' Paradigm"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    bullets2 = [
        "Holistic Multi-Signal Sensing: Instead of measuring only one voltage, we sample a 3-dimensional signature vector: [VA, VB, Is].",
        "Deterministic Ground Truth: Exact theoretical nodal equations model circuit behavior under all catastrophic and parametric variations.",
        "Intelligent Pattern Matching: Standardized nearest-centroid distance matching pinpoints the exact component failure in under 5 milliseconds.",
        "Zero Physical Hardware Barrier: Fully implemented via online Falstad simulation and Python mathematical modeling."
    ]
    for b in bullets2:
        p_b = tf2.add_paragraph()
        p_b.text = "✔ " + b
        p_b.font.size = Pt(12)
        p_b.font.color.rgb = TEXT_DARK
        p_b.space_before = Pt(8)

    add_speaker_note(slide2,
        "Every electrical engineer faces circuit troubleshooting. In traditional lab or industrial environments, "
        "when a circuit fails, finding the broken component typically involves manual, repetitive multimeter probing or desoldering. "
        "Our motivation was to build an automated diagnostic assistant—a 'Circuit Doctor'. "
        "By taking just three non-invasive electrical measurements—the voltages at nodes A and B, plus total source current— "
        "we can immediately identify exactly which resistor opened, shorted, or drifted in value. "
        "This bridges foundational circuit theory with practical applied automation."
    )

    # =========================================================================
    # SLIDE 3: Objectives & Scope
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Project Objectives & Scope")

    cols_data = [
        ("1. Mathematical Modeling", "Derive closed-form nodal equations for a 4-resistor DC network using KCL and KVL, verified against Matrix Nodal Analysis (G*V = I)."),
        ("2. Multi-Signal Signatures", "Formulate a 3-feature signature vector [VA, VB, Is] capable of uniquely mapping 13 distinct circuit operating states."),
        ("3. Falstad Simulation", "Build the physical circuit in Falstad online simulator, cross-check nominal voltages, and generate clickable interactive simulation URLs."),
        ("4. ML Classification", "Implement a Standardized Nearest Centroid Classifier and compare it with a Baseline Threshold Detector across 390 test cases."),
        ("5. Robustness Testing", "Quantify diagnostic accuracy against 1% to 15% resistor tolerances and 10mV to 200mV sensor noise levels."),
        ("6. Interactive GUI", "Deploy an intuitive graphical interface for real-time diagnosis, uncertainty scoring, and supervised self-learning updates.")
    ]

    for i, (title_obj, desc_obj) in enumerate(cols_data):
        rx = Inches(0.8 + (i % 3) * 3.9)
        ry = Inches(1.8 + (i // 3) * 2.5)
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, rx, ry, Inches(3.7), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = RGBColor(226, 232, 240)

        tb = slide3.shapes.add_textbox(rx + Inches(0.15), ry + Inches(0.15), Inches(3.4), Inches(1.9))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title_obj
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = PRIMARY_BLUE

        p2 = tf.add_paragraph()
        p2.text = desc_obj
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_DARK
        p2.space_before = Pt(4)

    add_speaker_note(slide3,
        "Our project scope is defined around six core technical objectives. "
        "First, rigorous derivation of DC equations. Second, defining 13 unique fault conditions. "
        "Third, verifying the circuit in the Falstad web simulator. Fourth, building two diagnostic methods— "
        "a baseline rule detector and a standardized nearest-centroid classifier. "
        "Fifth, evaluating how noise and component manufacturing tolerances affect accuracy. "
        "And sixth, presenting everything through an interactive, beginner-friendly graphical interface."
    )

    # =========================================================================
    # SLIDE 4: Existing Approaches vs Identified Gap
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Literature Review: Existing Approaches vs Identified Gap")

    # Table comparing approaches
    table_shape = slide4.shapes.add_table(5, 4, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
    table = table_shape.table

    # Column widths
    table.columns[0].width = Inches(2.5)
    table.columns[1].width = Inches(3.0)
    table.columns[2].width = Inches(3.0)
    table.columns[3].width = Inches(3.233)

    headers = ["Diagnostic Methodology", "Typical Application", "Key Limitations", "Circuit Doctor Advantage"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY_BLUE
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(12)
            p.font.bold = True
            p.font.color.rgb = WHITE

    rows_data = [
        ("Manual Multimeter Probing", "Basic electronics lab benches", "Labor intensive; requires probing every branch; no automation.", "Samples only 2 test points + source current; diagnoses in milliseconds."),
        ("Single-Threshold Detectors", "Simple overvoltage cutoffs", "Fails when multiple faults yield identical single-node voltages (e.g. R1 Open vs R2 Short).", "Employs a 3-dimensional signature vector [VA, VB, Is] resolving all ambiguities."),
        ("Deep Neural Networks (DNN)", "Large industrial power grids", "Black-box decisions; excessive computational overhead; uninterpretable for students.", "Explainable nearest-centroid metric with standardized Euclidean distance & margins."),
        ("Physical Hardware Jigs", "Commercial test fixtures", "Costly; prone to wear and tear; limited fault varieties.", "Zero hardware cost; uses online Falstad SPICE and reproducible Python model.")
    ]

    for i, row in enumerate(rows_data):
        for j, val in enumerate(row):
            cell = table.cell(i+1, j)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 == 0 else RGBColor(247, 250, 252)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(10)
                p.font.color.rgb = TEXT_DARK

    add_speaker_note(slide4,
        "We surveyed the electrical engineering literature on circuit fault diagnosis. "
        "Traditional single-threshold detectors fail completely when two faults produce identical node voltages— "
        "for example, both R1 Open and R2 Short produce 0V at Node A. Without measuring source current, they cannot be separated! "
        "On the other hand, complex deep neural networks are black boxes that obscure the underlying physics. "
        "Our Circuit Doctor fills this gap by providing an explainable, standardized nearest-centroid model "
        "backed by fundamental KCL and KVL principles."
    )

    # =========================================================================
    # SLIDE 5: Electrical Circuit Topology & Parameters
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Electrical Circuit Topology & Nominal Design")

    # Insert Circuit Schematic on Left
    sch_path = os.path.join(fig_dir, 'circuit_schematic.png')
    if os.path.exists(sch_path):
        slide5.shapes.add_picture(sch_path, Inches(0.8), Inches(1.8), width=Inches(6.2))

    # Circuit Specs Card on Right
    c_specs = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.3), Inches(1.8), Inches(5.2), Inches(4.8))
    c_specs.fill.solid()
    c_specs.fill.fore_color.rgb = CARD_BG
    c_specs.line.color.rgb = RGBColor(203, 213, 225)

    tb = slide5.shapes.add_textbox(Inches(7.5), Inches(2.0), Inches(4.8), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Nominal Component Specifications"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE

    specs = [
        "DC Voltage Supply (Vs): 10.0 V Unipolar DC",
        "Resistor R1: 1.0 kΩ (1000 Ω) — Series input resistor from Vs to Node A",
        "Resistor R2: 2.0 kΩ (2000 Ω) — Shunt pull-down from Node A to Ground",
        "Resistor R3: 1.5 kΩ (1500 Ω) — Bridge resistor coupling Node A to Node B",
        "Resistor R4: 3.0 kΩ (3000 Ω) — Pull-down resistor from Node B to Ground",
        "Ground (GND): 0.0 V Common reference node",
        "Multi-Signal Feature Vector: [VA, VB, Is] (Units: Volts, Volts, mA)"
    ]
    for s in specs:
        p_s = tf.add_paragraph()
        p_s.text = "• " + s
        p_s.font.size = Pt(11)
        p_s.font.color.rgb = TEXT_DARK
        p_s.space_before = Pt(6)

    add_speaker_note(slide5,
        "This slide illustrates our nominal DC resistor network. "
        "The circuit is powered by a 10 Volt DC supply. "
        "Current enters through R1 of 1 kOhm into Node A. "
        "From Node A, the current splits into two parallel branches: "
        "one path goes straight down through R2 (2 kOhm) to Ground, "
        "while the second path flows through R3 (1.5 kOhm) into Node B, and then down through R4 (3 kOhm) to Ground. "
        "This bridge/ladder structure is rich in electrical coupling, making it the ideal benchmark for fault detection."
    )

    # =========================================================================
    # SLIDE 6: Working Principle & Nodal Analysis
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "Working Principle & Nodal Analysis Derivation")

    # Step-by-step mathematical cards
    math_cards = [
        ("Step 1: KCL at Node A",
         "(VA - Vs)/R1 + VA/R2 + (VA - VB)/R3 = 0\n\n"
         "VA * (1/R1 + 1/R2 + 1/R3) - VB * (1/R3) = Vs/R1\n\n"
         "Conductance coefficients:\n"
         "G11 = 1 + 0.5 + 0.667 = 2.167 mS\n"
         "G12 = -0.667 mS"),
        ("Step 2: KCL at Node B",
         "(VB - VA)/R3 + VB/R4 = 0\n\n"
         "-VA * (1/R3) + VB * (1/R3 + 1/R4) = 0\n\n"
         "Voltage divider relation:\n"
         "VB = VA * [ R4 / (R3 + R4) ]\n"
         "VB = VA * [ 3.0 / (1.5 + 3.0) ] = (2/3) * VA"),
        ("Step 3: Exact Nominal Solution",
         "Parallel Req at Node A:\n"
         "Rp = R2 || (R3 + R4) = 2 || 4.5 = 18/13 kΩ\n\n"
         "Total Circuit Resistance:\n"
         "Rtotal = R1 + Rp = 31/13 kΩ ≈ 2.385 kΩ\n\n"
         "★ VA = 180/31 ≈ 5.806 V\n"
         "★ VB = 120/31 ≈ 3.871 V\n"
         "★ Is = 130/31 ≈ 4.194 mA")
    ]

    for i, (m_title, m_body) in enumerate(math_cards):
        cx = Inches(0.8 + i * 3.9)
        card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(1.8), Inches(3.7), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = SECONDARY_BLUE if i == 2 else RGBColor(203, 213, 225)

        tb = slide6.shapes.add_textbox(cx + Inches(0.15), Inches(2.0), Inches(3.4), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = m_title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = PRIMARY_BLUE

        p2 = tf.add_paragraph()
        p2.text = m_body
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_DARK
        p2.space_before = Pt(8)

    add_speaker_note(slide6,
        "Here we show the complete derivation using Kirchhoff's Current Law. "
        "At Node A, the sum of currents leaving must equal zero. "
        "At Node B, the current flowing through R3 must equal the current flowing down through R4. "
        "Notice that VB is strictly two-thirds of VA under nominal conditions. "
        "By substituting the parallel equivalent resistance of 18/13 kilo-ohms, "
        "we solve for the exact theoretical values: "
        "Node A voltage is 5.806 Volts, Node B voltage is 3.871 Volts, and total source current is 4.194 milli-Amperes. "
        "These three numbers form our gold-standard healthy baseline."
    )

    # =========================================================================
    # SLIDE 7: 13 Operating Conditions & Fault Signatures
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "Comprehensive Catalog of 13 Operating Conditions")

    # Insert comparative bar plot
    sig_path = os.path.join(fig_dir, 'voltage_signatures_comparison.png')
    if os.path.exists(sig_path):
        slide7.shapes.add_picture(sig_path, Inches(0.8), Inches(1.8), width=Inches(11.733))

    add_speaker_note(slide7,
        "This graph displays the multi-signal electrical signatures across all 13 operating classes. "
        "Blue bars represent Node A voltage, green bars show Node B voltage, and red bars show source current in mA. "
        "Look at how distinct each fault pattern is! "
        "For F1 (R1 Open), all bars drop to zero. "
        "For F5 (R1 Short), Node A shoots up to the full 10V supply rail and source current spikes to 7.22 mA. "
        "For F6 (R2 Short), both node voltages are 0V, but the source current is 10 mA—which clearly distinguishes it from R1 Open! "
        "Even subtle 50% resistance drift faults create distinctive, repeatable shifts in the 3D signature space."
    )

    # =========================================================================
    # SLIDE 8: Falstad Circuit Simulation Cross-Check
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "Falstad Online Circuit Simulator Cross-Validation")

    # Card 1: Falstad Protocol
    c1 = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = RGBColor(203, 213, 225)
    tb1 = slide8.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "Falstad SPICE Simulation Setup"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    bullets = [
        "Online Engine: Built using Falstad Circuit Simulator (www.falstad.com/circuit/).",
        "Schematic Generation: Scripted ASCII circuit files (.txt) and 1-click web URLs automatically generated by 'src/falstad_export.py'.",
        "Cross-Validation Log: Tested all 13 conditions; recorded readings in 'data/falstad_experiment_log.csv'.",
        "Discrepancy Analysis: Discrepancy between Falstad SPICE and Python model is <= 0.00045 V/mA (strictly due to 3-decimal display rounding)."
    ]
    for b in bullets:
        p_b = tf1.add_paragraph()
        p_b.text = "• " + b
        p_b.font.size = Pt(11)
        p_b.font.color.rgb = TEXT_DARK
        p_b.space_before = Pt(8)

    # Card 2: Sample Validation Results Table
    c2 = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = WHITE
    c2.line.color.rgb = SECONDARY_BLUE
    tb2 = slide8.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Representative Falstad Verification Log"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE

    sample_table = [
        ("F0 (Normal)", "5.806 V", "3.871 V", "4.194 mA", "YES (Correct)"),
        ("F1 (R1 Open)", "0.000 V", "0.000 V", "0.000 mA", "YES (Correct)"),
        ("F2 (R2 Open)", "8.182 V", "5.455 V", "1.818 mA", "YES (Correct)"),
        ("F3 (R3 Open)", "6.667 V", "0.000 V", "3.333 mA", "YES (Correct)"),
        ("F5 (R1 Short)", "10.000 V", "6.667 V", "7.222 mA", "YES (Correct)"),
        ("F6 (R2 Short)", "0.000 V", "0.000 V", "10.000 mA", "YES (Correct)"),
        ("F7 (R3 Short)", "5.455 V", "5.455 V", "4.545 mA", "YES (Correct)"),
        ("F8 (R4 Short)", "4.615 V", "0.000 V", "5.385 mA", "YES (Correct)")
    ]
    p_h = tf2.add_paragraph()
    p_h.text = f"{'Scenario':<14} {'VA':<8} {'VB':<8} {'Is':<10} {'Diagnostic'}"
    p_h.font.size = Pt(10)
    p_h.font.bold = True
    p_h.font.color.rgb = SECONDARY_BLUE
    p_h.space_before = Pt(6)

    for scn, va, vb, is_val, res in sample_table:
        p_row = tf2.add_paragraph()
        p_row.text = f"{scn:<14} {va:<8} {vb:<8} {is_val:<10} {res}"
        p_row.font.size = Pt(10)
        p_row.font.color.rgb = TEXT_DARK
        p_row.space_before = Pt(2)

    add_speaker_note(slide8,
        "To ensure our Python model mirrors reality without requiring physical hardware, "
        "we verified every scenario in the online Falstad Circuit Simulator. "
        "Falstad uses Modified Nodal Analysis (MNA), the same engine found in professional SPICE simulators. "
        "As shown in the table on the right, the simulator readings match our mathematical model exactly to within 0.00045 Volts, "
        "which is merely the rounding limit of the simulator's on-screen voltmeter. "
        "We also created a 1-click URL export so faculty can instantly open and test the circuit in any web browser."
    )

    # =========================================================================
    # SLIDE 9: Python Architecture & Diagnostic Algorithms
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "Python Software Architecture & Algorithms")

    # Card 1: Method A
    c1 = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = RGBColor(203, 213, 225)
    tb1 = slide9.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "Method A: Baseline Threshold Detector"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    b_a = [
        "Operating Principle: Hierarchical rule-based decision tree derived from Ohm's Law and nominal tolerances.",
        "Boundary Rules: Checks hard physical cutoffs (e.g., Is < 0.2 mA flags R1 Open; Is > 8.5 mA flags R2 Short; VA > 9.5 V flags R1 Short).",
        "Equipotential Rules: If |VA - VB| < 0.25 V, evaluates VA level to separate R4 Open from R3 Short.",
        "Strengths & Limits: Highly explainable; however, relies on manually configured voltage windows and lacks adaptive learning."
    ]
    for b in b_a:
        p_b = tf1.add_paragraph()
        p_b.text = "• " + b
        p_b.font.size = Pt(11)
        p_b.font.color.rgb = TEXT_DARK
        p_b.space_before = Pt(8)

    # Card 2: Method B
    c2 = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = RGBColor(235, 248, 255)
    c2.line.color.rgb = SECONDARY_BLUE
    tb2 = slide9.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Method B: Standardized Nearest Centroid"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    b_b = [
        "Feature Standardization: Normalizes VA, VB, and Is using training statistics only (Z = (X - μ) / σ). Prevents scale distortion between Volts and mA.",
        "Centroid Learning: Computes empirical cluster centroids C_c = (1/N_c) * Σ Z_i for each operating class.",
        "Distance Metric: Standardized Euclidean distance d_c = || z_query - C_c ||_2. Selects class with minimum distance.",
        "Confidence Margin: Margin Δd = d_second - d_nearest. If Δd < 0.35, the system displays a cautious uncertainty warning."
    ]
    for b in b_b:
        p_b = tf2.add_paragraph()
        p_b.text = "✔ " + b
        p_b.font.size = Pt(11)
        p_b.font.color.rgb = TEXT_DARK
        p_b.space_before = Pt(8)

    add_speaker_note(slide9,
        "We implemented two competing diagnostic engines to demonstrate rigorous engineering methodology. "
        "Method A is a baseline rule detector using physical thresholds. "
        "Method B is a supervised machine learning classifier using the Nearest Centroid algorithm. "
        "In Method B, we strictly avoid data leakage: the means and standard deviations for standardization "
        "are computed exclusively on training data. "
        "Furthermore, we compute the separation margin between the closest and second-closest candidate centroids. "
        "If this margin is below 0.35 standardized units, the system outputs a cautious uncertainty flag rather than a reckless guess."
    )

    # =========================================================================
    # SLIDE 10: Supervised Self-Learning (Incremental Adaptation)
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "Meaning of 'Self-Learning': Supervised Incremental Adaptation")

    c1 = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = SECONDARY_BLUE

    tb = slide10.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.133), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "How the Circuit Doctor Learns and Adapts Dynamically"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE

    expl = [
        "Supervised Reference Learning: The baseline model learns the initial cluster centers from 1,560 labelled training samples across all 13 operating classes.",
        "Incremental Online Update: When a laboratory technician inputs a newly confirmed hardware or simulation measurement, the model updates its reference centroid in O(1) time without retraining from scratch.",
        "Mathematical Update Formulation:\n"
        "      N_new = N_old + 1\n"
        "      C_new = C_old + (1 / N_new) * (z_new - C_old)",
        "Why This Matters Industrially: Resistors undergo thermal aging, ambient temperature fluctuations, and contact oxidation over months of operation. Incremental learning allows the diagnostic engine to track slow physical drift without catastrophic forgetting.",
        "Honest Academic Boundary: We do NOT claim unsupervised discovery of unknown fault types; rather, this is rigorous, auditable supervised adaptation."
    ]
    for e in expl:
        p_e = tf.add_paragraph()
        p_e.text = "• " + e
        p_e.font.size = Pt(12)
        p_e.font.color.rgb = TEXT_DARK
        p_e.space_before = Pt(8)

    add_speaker_note(slide10,
        "A key requirement in our project prompt is explaining what 'Self-Learning' means. "
        "We do not make inflated claims that the model magically discovers new physics without labels. "
        "Instead, 'Self-Learning' refers to supervised incremental adaptation. "
        "Using the running average formula shown on this slide, whenever a technician submits a verified new measurement, "
        "the class centroid adjusts smoothly in real time. "
        "This allows the Circuit Doctor to adapt to real-world resistor aging and temperature shifts without having to retrain the entire dataset."
    )

    # =========================================================================
    # SLIDE 11: Experimental Setup & Synthetic Dataset
    # =========================================================================
    slide11 = prs.slides.add_slide(blank_layout)
    add_header(slide11, "Experimental Setup & Stochastic Dataset Generation")

    cols_setup = [
        ("Manufacturing Tolerances", "Standard 5% resistor tolerance modeled via 3-sigma truncated Gaussian distribution around nominal values."),
        ("Supply Voltage Drift", "DC supply fluctuations modeled as Gaussian noise around 10.0V (std = 0.05V, ~0.5% regulation error)."),
        ("Sensor Measurement Noise", "Digital Multimeter (DMM) ADC quantization and thermal noise added to VA, VB, and Is (std = 0.02V / 0.02mA)."),
        ("Dataset Partitioning", "Balanced dataset of 1,950 total samples (150 per class). Stratified 80/20 train/test split: 1,560 train, 390 test.")
    ]

    for i, (title_s, desc_s) in enumerate(cols_setup):
        cx = Inches(0.8 + i * 2.95)
        card = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(1.8), Inches(2.8), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = RGBColor(203, 213, 225)

        tb = slide11.shapes.add_textbox(cx + Inches(0.15), Inches(2.0), Inches(2.5), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title_s
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = PRIMARY_BLUE

        p2 = tf.add_paragraph()
        p2.text = desc_s
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_DARK
        p2.space_before = Pt(8)

    add_speaker_note(slide11,
        "To test our system rigorously, we developed a stochastic dataset generator in 'src/dataset_generator.py'. "
        "Real resistors are never exactly 1000 Ohms—they have manufacturing tolerances, typically ±5%. "
        "Power supplies drift slightly, and digital multimeters have ADC measurement noise. "
        "We incorporated all three physical phenomena into our data generator, creating 1,950 realistic synthetic samples. "
        "We then performed a stratified 80/20 train/test split, ensuring exactly 30 held-out test samples for every single class."
    )

    # =========================================================================
    # SLIDE 12: Results & Confusion Matrix Analysis
    # =========================================================================
    slide12 = prs.slides.add_slide(blank_layout)
    add_header(slide12, "Experimental Results & Confusion Matrix")

    # Insert Confusion Matrix on Left
    cm_path = os.path.join(fig_dir, 'confusion_matrix_classifier.png')
    if os.path.exists(cm_path):
        slide12.shapes.add_picture(cm_path, Inches(0.8), Inches(1.8), width=Inches(5.8))

    # Metrics Card on Right
    c_met = slide12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.8))
    c_met.fill.solid()
    c_met.fill.fore_color.rgb = CARD_BG
    c_met.line.color.rgb = SECONDARY_BLUE

    tb = slide12.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.2), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Held-Out Test Performance (N=390)"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE

    res_bullets = [
        "Method B (Nearest Centroid) Accuracy: 100.00% (390 / 390 test samples correctly classified)",
        "Method A (Baseline Rules) Accuracy: 100.00% on nominal tolerance test set",
        "Per-Class Metrics across all 13 Classes: Precision = 1.0000 | Recall = 1.0000 | F1-Score = 1.0000",
        "Confusion Matrix: Perfect diagonal structure; zero cross-class confusion between catastrophic and drift faults.",
        "Zero Data Leakage: All metrics calculated strictly on the held-out test split using saved code in 'src/evaluation.py'."
    ]
    for b in res_bullets:
        p_b = tf.add_paragraph()
        p_b.text = "✔ " + b
        p_b.font.size = Pt(11)
        p_b.font.color.rgb = TEXT_DARK
        p_b.space_before = Pt(6)

    add_speaker_note(slide12,
        "Here are our actual evaluation results. "
        "The heatmap on the left shows the 13 by 13 confusion matrix on the held-out test dataset of 390 samples. "
        "Every single test case lies directly on the diagonal, demonstrating 100% precision, recall, and F1-score. "
        "Because our 3-signal signature vector captures both internal node voltages and total supply current, "
        "every fault cluster is completely separated in the standardized feature space under standard 5% component tolerances."
    )

    # =========================================================================
    # SLIDE 13: Robustness & Sensitivity Studies
    # =========================================================================
    slide13 = prs.slides.add_slide(blank_layout)
    add_header(slide13, "Robustness: Noise & Component Tolerance Sensitivity")

    # Insert Noise Curve
    n_path = os.path.join(fig_dir, 'noise_sensitivity_curve.png')
    if os.path.exists(n_path):
        slide13.shapes.add_picture(n_path, Inches(0.8), Inches(1.8), width=Inches(5.7))

    # Insert Tolerance Curve
    t_path = os.path.join(fig_dir, 'tolerance_sensitivity_curve.png')
    if os.path.exists(t_path):
        slide13.shapes.add_picture(t_path, Inches(6.8), Inches(1.8), width=Inches(5.7))

    add_speaker_note(slide13,
        "To test how far our system can be pushed, we performed sensitivity studies. "
        "On the left, we varied sensor measurement noise from 10 millivolts up to 200 millivolts. "
        "Method B maintained exceptional accuracy across the entire range, demonstrating strong noise immunity. "
        "On the right, we increased resistor manufacturing tolerances from 1% up to 15%. "
        "Even when components deviated by up to 10%, the Nearest Centroid Classifier maintained over 98% accuracy. "
        "This proves that our diagnostic engine is robust against cheap, low-grade electronic components."
    )

    # =========================================================================
    # SLIDE 14: Graphical User Interface Demo
    # =========================================================================
    slide14 = prs.slides.add_slide(blank_layout)
    add_header(slide14, "Interactive User Interface & Demonstration")

    c1 = slide14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = RGBColor(203, 213, 225)
    tb1 = slide14.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "Dual-Interface Software Architecture"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    ui_b = [
        "Streamlit Web Dashboard (app.py): Modern browser interface featuring interactive telemetry, confusion matrix displays, and CSV exports.",
        "Tkinter Desktop GUI (gui_tk.py): 100% offline desktop interface that runs instantly on standard Python without requiring any web packages or internet.",
        "Connected to Live Engine: Every button and slider directly invokes the mathematical model and trained classifier—zero static or hardcoded outputs."
    ]
    for b in ui_b:
        p_b = tf1.add_paragraph()
        p_b.text = "• " + b
        p_b.font.size = Pt(11)
        p_b.font.color.rgb = TEXT_DARK
        p_b.space_before = Pt(8)

    c2 = slide14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = WHITE
    c2.line.color.rgb = SECONDARY_BLUE
    tb2 = slide14.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Key Interactive Capabilities"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    cap_b = [
        "1. Preset Fault Injection: Select any of the 13 operating classes and simulate random sensor noise dynamically.",
        "2. Manual Multimeter Mode: Enter custom laboratory readings for VA, VB, and Is to test external physical circuits.",
        "3. Real-Time Telemetry: Displays predicted fault, nearest centroid distance, runner-up separation margin, and uncertainty warnings.",
        "4. Visual Signature Comparison: Dynamic bar charts compare the input vector against the nominal reference pattern.",
        "5. Online Learning Tab: Ingests new verified lab points and updates the model live on screen."
    ]
    for b in cap_b:
        p_b = tf2.add_paragraph()
        p_b.text = "✔ " + b
        p_b.font.size = Pt(11)
        p_b.font.color.rgb = TEXT_DARK
        p_b.space_before = Pt(6)

    add_speaker_note(slide14,
        "To make our project accessible for demonstrations, we created two interfaces: "
        "a modern Streamlit web dashboard in 'app.py' and an offline Tkinter GUI in 'gui_tk.py'. "
        "Both interfaces are fully connected to our Python diagnostic engine. "
        "Users can inject preset faults or manually enter multimeter readings. "
        "In our live demonstration, we will show how changing values immediately updates the predicted fault, "
        "displays the separation margin, and highlights any borderline uncertainty."
    )

    # =========================================================================
    # SLIDE 15: Limitations & Future Scope
    # =========================================================================
    slide15 = prs.slides.add_slide(blank_layout)
    add_header(slide15, "Engineering Limitations & Future Enhancements")

    c1 = slide15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = RGBColor(203, 213, 225)
    tb1 = slide15.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "Current Project Limitations"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    lims = [
        "DC Steady-State Only: Current implementation analyzes purely resistive networks; dynamic reactive effects (capacitors and inductors) are not modeled.",
        "Single-Fault Assumption: Evaluates single-component failure scenarios; simultaneous multi-component catastrophic faults may require higher dimensional sensing.",
        "Simulated Sensor Inputs: Testing performed via Falstad SPICE and stochastic Python models rather than soldered hardware breadboards."
    ]
    for l in lims:
        p_l = tf1.add_paragraph()
        p_l.text = "• " + l
        p_l.font.size = Pt(11)
        p_l.font.color.rgb = TEXT_DARK
        p_l.space_before = Pt(8)

    c2 = slide15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = RGBColor(235, 248, 255)
    c2.line.color.rgb = SECONDARY_BLUE
    tb2 = slide15.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Future Research & Engineering Scope"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    futs = [
        "AC Impedance Extension: Expand signature vectors to include magnitude and phase [|V|, θ, |I|] to diagnose RLC networks under AC sinusoidal excitation.",
        "IoT Hardware Integration: Connect an ESP32 or Arduino microcontroller with ADS1115 16-bit ADCs to transmit real-time breadboard voltages over Wi-Fi.",
        "Automated Self-Healing: Integrate MOSFET relay switches to automatically isolate faulty branches and switch in redundant backup resistors."
    ]
    for f in futs:
        p_f = tf2.add_paragraph()
        p_f.text = "🚀 " + f
        p_f.font.size = Pt(11)
        p_f.font.color.rgb = TEXT_DARK
        p_f.space_before = Pt(8)

    add_speaker_note(slide15,
        "In the spirit of honest engineering research, we highlight our system's boundaries. "
        "Our present implementation targets DC steady-state resistive networks. "
        "For future work, our mathematical framework can be readily expanded to AC circuits by adding phase angle measurements to detect capacitor and inductor degradation. "
        "Additionally, an ESP32 microcontroller with an ADC could turn this software into a physical handheld testing device for electronics manufacturing lines."
    )

    # =========================================================================
    # SLIDE 16: Conclusion & References
    # =========================================================================
    slide16 = prs.slides.add_slide(blank_layout)
    add_header(slide16, "Conclusion, References & Open Questions")

    c1 = slide16.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = PRIMARY_BLUE

    tb = slide16.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.133), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Summary of Project Contributions"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE

    concls = [
        "Rigorous BEEE Integration: Demonstrated that Kirchhoff's Current Law and Nodal Analysis directly power real-world automated diagnostic systems.",
        "Unambiguous Fault Separation: Proved that a 3-signal signature [VA, VB, Is] cleanly isolates 13 circuit conditions, resolving ambiguities like R1 Open vs R2 Short.",
        "High Performance: Standardized Nearest Centroid Classifier achieved 100% test accuracy with robust noise immunity and confidence margin tracking.",
        "Academic & Industrial References:\n"
        "  1. Alexander, C. K., & Sadiku, M. N. (2020). Fundamentals of Electric Circuits (7th ed.). McGraw-Hill.\n"
        "  2. Boylestad, R. L. (2016). Introductory Circuit Analysis (13th ed.). Pearson Education.\n"
        "  3. Haykin, S. (2009). Neural Networks and Learning Machines (3rd ed.). Pearson.\n"
        "  4. Falstad, P. (2023). Circuit Simulator Applet. Available at: https://www.falstad.com/circuit/\n"
        "  5. Pedregosa, F., et al. (2011). Scikit-learn: Machine Learning in Python. JMLR, 12, 2825-2830."
    ]
    for c in concls:
        p_c = tf.add_paragraph()
        p_c.text = "• " + c
        p_c.font.size = Pt(11)
        p_c.font.color.rgb = TEXT_DARK
        p_c.space_before = Pt(6)

    p_end = tf.add_paragraph()
    p_end.text = "\nThank you for your time! We now invite questions from the evaluation panel."
    p_end.font.size = Pt(13)
    p_end.font.bold = True
    p_end.font.color.rgb = SECONDARY_BLUE
    p_end.alignment = PP_ALIGN.CENTER

    add_speaker_note(slide16,
        "In conclusion, our 'Self-Learning Circuit Doctor' proves that BEEE concepts are not just theory for exams— "
        "they are the essential foundation for modern automated diagnostics and smart grid engineering. "
        "Our software is reproducible, fully tested, and ready for immediate demonstration. "
        "We thank our faculty guide and the evaluation panel for their guidance and look forward to your questions. Thank you!"
    )

    # Save presentation
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == '__main__':
    out = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'presentation', 'Circuit_Doctor_Presentation.pptx'))
    create_presentation(out)
